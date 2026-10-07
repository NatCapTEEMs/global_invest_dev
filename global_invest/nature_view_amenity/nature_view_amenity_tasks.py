"""Nature-view-amenity GEP tasks: the hotel-premium valuation on the staged CAS workbooks.

This layer owns every file read and write. The science it calls lives in
nature_view_amenity_functions, which never opens a file.

The room sample, the comparable groups, the country mapping and the UN Tourism room inventory
are the CAS team's staged submission, taken as given; the owner's results workbook is carried
beside the result as the replication anchor and the run log reports the per-country agreement.
The room prices are scraped at the submission's own dates and applied to 2019 room-nights,
which is the owner's convention, kept.
"""
import os

import hazelbean as hb
import numpy as np
import pandas as pd

from global_invest import utilities
from global_invest.nature_view_amenity import nature_view_amenity_functions as nv

LANDSCAPE_CLASSES = ('Natural scenic view', 'Non-scenic', 'Built-environment view')


def publish_inputs(p):
    """Every GEP task's first line: the nature_view_amenity es_config row and the parameter
    rows from es_parameters (defaults layer -- a caller-set value prevails), the shared country
    references and the results registry."""
    utilities.hydrate_es_config(p, 'nature_view_amenity', log=hb.log)
    utilities.hydrate_es_parameters(p, 'nature_view_amenity', log=hb.log)
    utilities.initialize_country_paths(p, simplified='30sec')
    if not hasattr(p, 'results'):
        p.results = {}
    return p


def coordinate_countries(p):
    """The coordinate-to-country mapping, read once from the raw hotel master and cached.

    The master is a quarter-gigabyte workbook consulted for two columns, so the extraction
    is cached as a CSV in this task's directory.
    """
    publish_inputs(p)
    p.nature_view_coordinate_country_path = os.path.join(p.cur_dir, 'coordinate_iso2.csv')
    if not p.run_this:
        return
    if not hb.path_exists(p.nature_view_coordinate_country_path):
        from openpyxl import load_workbook
        wb = load_workbook(str(p.get_path(p.nature_view_amenity_raw_master_input_path)),
                           read_only=True, data_only=True)
        ws = wb['Nearest_Hotel']
        rows = ws.iter_rows(values_only=True)
        header = {name: i for i, name in enumerate(next(rows))}
        pairs = {}
        for row in rows:
            coordinate_id = str(row[header['coordinate_id']] or '').strip()
            country_code = str(row[header['country_code']] or '').strip().upper()
            if coordinate_id and country_code:
                pairs[coordinate_id] = country_code
        wb.close()
        pd.DataFrame(sorted(pairs.items()), columns=['coordinate_id', 'iso2']).to_csv(
            p.nature_view_coordinate_country_path, index=False, encoding='utf-8-sig')
        hb.log('  coordinate-to-country mapping cached for %d coordinates' % len(pairs))
    return True


def gep_calculation(p):
    """GEP valuation for the nature-view amenity: hotel rooms times occupancy times 365 times
    the scenic-room share times the within-hotel unit price difference, per country, anchored
    against the owner's results workbook."""
    publish_inputs(p)
    service_results, already_done = utilities.begin_gep_calculation(p, 'nature_view_amenity')
    if already_done:
        return

    # keep_default_na: Namibia's alpha-2 code is the literal string NA, which pandas would
    # otherwise read as missing and silently drop the country.
    mapping = pd.read_excel(str(p.get_path(p.nature_view_amenity_country_mapping_input_path)),
                            keep_default_na=False, na_values=[''])
    iso2_to_iso3 = dict(zip(mapping['ISO Alpha-2'].astype(str).str.strip().str.upper(),
                            mapping['ISO Alpha-3'].astype(str).str.strip().str.upper()))
    coord = pd.read_csv(p.nature_view_coordinate_country_path, dtype=str,
                        keep_default_na=False, na_values=[''])
    coord['iso3'] = coord['iso2'].map(iso2_to_iso3)

    rooms = pd.read_excel(str(p.get_path(p.nature_view_amenity_comparable_groups_input_path)),
                          sheet_name='Room_Inventory_Summary',
                          usecols=['coordinate_id', 'avg_all_inclusive_price',
                                   'available_room_count', 'landscape_class',
                                   'comparable_group_id', 'comparison_role'])
    rooms['coordinate_id'] = rooms['coordinate_id'].astype(str).str.strip()
    rooms = rooms.merge(coord[['coordinate_id', 'iso3']], on='coordinate_id', how='left')
    rooms['price'] = pd.to_numeric(rooms['avg_all_inclusive_price'], errors='coerce')
    rooms['availability'] = pd.to_numeric(rooms['available_room_count'], errors='coerce')

    # the scenic-room share: every valid room record, natural availability over the total
    share_rows = rooms[rooms['landscape_class'].isin(LANDSCAPE_CLASSES)
                       & (rooms['price'] > 0) & (rooms['availability'] > 0)
                       & rooms['iso3'].notna()]
    n_invalid = int((rooms['landscape_class'].isin(LANDSCAPE_CLASSES)
                     & ~((rooms['price'] > 0) & (rooms['availability'] > 0))).sum())
    shares = share_rows.groupby('iso3').apply(
        lambda g: pd.Series({
            'natural_availability': g.loc[g['landscape_class'] == 'Natural scenic view',
                                          'availability'].sum(),
            'total_availability': g['availability'].sum()}), include_groups=False).reset_index()
    shares['scenic_room_share'] = shares['natural_availability'] / shares['total_availability']
    hb.log('  scenic shares from %d valid room records (%d excluded for missing price or '
           'availability)' % (len(share_rows), n_invalid))

    # the within-hotel premium: per comparable group, mean scenic and comparable prices
    grouped = rooms[rooms['comparable_group_id'].astype(str).str.strip().ne('')
                    & rooms['comparable_group_id'].notna()]
    natural = grouped[grouped['comparison_role'] == 'natural_target'].groupby('comparable_group_id').agg(
        scenic_price=('price', 'mean'), scenic_availability=('availability', 'sum'),
        iso3=('iso3', 'first'))
    control = grouped[grouped['comparison_role'] == 'non_scenic_control'].groupby(
        'comparable_group_id').agg(non_scenic_price=('price', 'mean'))
    groups = natural.join(control, how='inner').reset_index()
    groups = nv.summarize_groups(groups)
    groups['is_outlier'] = nv.flag_3iqr_outliers(groups['log_price_ratio'])
    upd = nv.country_unit_price_difference(groups)
    hb.log('  %d comparable groups, %d outliers, %d kept into the premium'
           % (len(groups), int(groups['is_outlier'].sum()),
              int(upd['n_groups'].sum())))

    un = pd.read_excel(str(p.get_path(p.nature_view_amenity_un_tourism_input_path)),
                       sheet_name='Country Data')
    un['iso3'] = un['ISO Alpha-3'].astype(str).str.strip().str.upper()
    occupancy = pd.to_numeric(un['Room Occupancy Rate 2019 (%)'], errors='coerce').dropna()
    mean_occupancy = float(occupancy.mean()) / 100.0
    hb.log('  uniform mean occupancy %.6f from %d reporting countries'
           % (mean_occupancy, len(occupancy)))

    df = un[['iso3', 'Hotel Rooms 2019']].rename(columns={'Hotel Rooms 2019': 'hotel_rooms'})
    df['hotel_rooms'] = pd.to_numeric(df['hotel_rooms'], errors='coerce')
    df = df.merge(shares[['iso3', 'scenic_room_share']], on='iso3', how='outer')
    df = df.merge(upd[['iso3', 'unit_price_difference']], on='iso3', how='outer')
    df = nv.nature_view_value(df, mean_occupancy)

    countries = utilities.collapse_countries_to_r250(p.df_countries)
    df_gep = countries.merge(df.rename(columns={'iso3': 'iso3_r250_label'}),
                             on='iso3_r250_label', how='left')
    unmapped = sorted(set(df['iso3'].dropna()) - set(countries['iso3_r250_label']))
    if unmapped:
        hb.log('  ⚠ %d submission codes are not on the account list and are dropped: %s'
               % (len(unmapped), ', '.join(unmapped[:8])))

    reference = pd.read_excel(str(p.get_path(p.nature_view_amenity_reference_path)))
    reference['iso3_r250_label'] = reference['Country or Area'].map(
        dict(zip(mapping['Country or Area'], mapping['ISO Alpha-3'])))
    df_gep = df_gep.merge(reference.rename(columns={
        'Annual Premium Value (USD)': 'nature_view_amenity_gep_reference'})[
        ['iso3_r250_label', 'nature_view_amenity_gep_reference']],
        on='iso3_r250_label', how='left')

    # the anchor holds at the owner's scraped price vintage; the published table is the
    # account's base-year dollars via the shared CPI series
    cpi = pd.read_csv(str(p.get_path(p.nature_view_amenity_us_cpi_input_path)))
    deflation = utilities.usd_deflation_factor(
        cpi, int(p.nature_view_amenity_price_year), int(p.gep_base_year))
    df_gep['nature_view_amenity_gep_owner_vintage'] = df_gep['nature_view_amenity_gep']
    df_gep['nature_view_amenity_gep'] = df_gep['nature_view_amenity_gep'] * deflation
    hb.log('  published value deflated from the %s scrape vintage to %s USD by the shared '
           'CPI factor %.6f: total %s'
           % (p.nature_view_amenity_price_year, p.gep_base_year, deflation,
              f"{df_gep['nature_view_amenity_gep'].sum():,.2f}"))
    df_gep['year'] = int(p.gep_base_year)
    utilities.write_gep_by_country(
        p, df_gep[utilities.published_country_columns(df_gep, 'nature_view_amenity')],
        service_results['gep_by_country_base_year'])
    gdf = hb.df_merge(p.gdf_countries_simplified, df_gep, how='outer',
                      left_on='ee_r264_id', right_on='ee_r264_id')
    gdf.to_file(service_results['gep_by_country_base_year'].replace('.csv', '.gpkg'), driver='GPKG')

    ours = df_gep['nature_view_amenity_gep_owner_vintage'].sum()
    ref = df_gep['nature_view_amenity_gep_reference'].sum()
    both = df_gep[df_gep['nature_view_amenity_gep_owner_vintage'].notna()
                  & df_gep['nature_view_amenity_gep_reference'].notna()]
    agree = (abs(both['nature_view_amenity_gep_owner_vintage'] - both['nature_view_amenity_gep_reference'])
             <= 0.01 + 1e-9 * both['nature_view_amenity_gep_reference'].abs()).sum()
    hb.log(f'At the owner price vintage: {ours:,.2f} '
           f'({int((df_gep["nature_view_amenity_gep"].fillna(0) > 0).sum())} countries); '
           f'owner reference {ref:,.2f}; countries agreeing to the cent: {agree} of {len(both)}')
    return True


def gep_result(p):
    """Render the results report(s). Shared implementation in utilities."""
    publish_inputs(p)
    utilities.render_service_results(p)
